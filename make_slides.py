"""Build slides.pptx (6 slides: Title, Problem, Method, Results x2, Takeaway).
Run after make_figures.py: python make_slides.py
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x4C, 0x72, 0xB0)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=20, color=NAVY, bold=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size), color, bold
        p.space_after = Pt(10)
    return tf


def header(slide, kicker, title):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    text(slide, 0.6, 0.35, 12, 0.4, [kicker.upper()], size=14, color=BLUE, bold=True)
    text(slide, 0.6, 0.75, 12, 0.9, [title], size=32, bold=True)


def bullets(slide, x, y, w, items, size=20):
    text(slide, x, y, w, 5, [f"•  {i}" for i in items], size=size, color=GREY)


# 1 Title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(s, 0.9, 2.2, 11.5, 2, ["Predicting Concrete Compressive Strength for Unseen Mixtures"],
     size=40, color=WHITE, bold=True)
text(s, 0.9, 3.9, 11.5, 1, ["Tree ensembles, ratio features, and the cost of random splitting"],
     size=22, color=RGBColor(0xC8, 0xD3, 0xE8))
text(s, 0.9, 5.6, 11.5, 0.6, ["Mohiminul Islam Fardin  ·  October 2026"], size=16,
     color=RGBColor(0xC8, 0xD3, 0xE8))

# 2 Problem
s = prs.slides.add_slide(BLANK)
header(s, "Problem", "Strength is confirmed only after 28 days")
bullets(s, 0.6, 1.9, 7.2, [
    "Predicting strength from the mix design would cut trial batches and support cost- and carbon-aware design",
    "Published ML models report R² above 0.95 on the UCI dataset (Feng et al., 2020)",
    "But most split the data at random, while each mixture is tested at several ages",
    "So the same recipe can appear in both training and test sets",
])
text(s, 8.4, 2.0, 4.4, 3.5, [
    "Research question",
    "How accurately can strength be predicted for mixtures the model has never seen, "
    "and do tree ensembles with ratio features beat a linear baseline?",
], size=19, color=NAVY)

# 3 Method
s = prs.slides.add_slide(BLANK)
header(s, "Method", "Grouped evaluation on 428 distinct mixtures")
bullets(s, 0.6, 1.9, 6.2, [
    "UCI dataset (Yeh, 1998): 1,030 tests → 996 records after removing duplicates and averaging replicates",
    "Features: 8 raw inputs vs 15 engineered (water/binder, water/cement, log-age, …)",
    "Models: Linear Regression, Random Forest, Gradient Boosting",
], size=19)
bullets(s, 7.0, 1.9, 5.8, [
    "Test set: 20% of mixtures held out whole (86 mixtures)",
    "Tuning: randomised search, 5-fold GroupKFold on training data",
    "Metrics: RMSE (primary), MAE, R²",
    "Leakage check: grouped vs random split over 10 seeds",
], size=19)

# 4 Results: model comparison
s = prs.slides.add_slide(BLANK)
header(s, "Results", "Gradient Boosting + ratio features: RMSE 4.40 MPa, R² 0.926")
s.shapes.add_picture("figures/fig1_test_rmse.png", Inches(0.5), Inches(1.8), height=Inches(5.2))
bullets(s, 9.4, 2.0, 3.6, [
    "Ensembles clearly beat the linear baseline",
    "Ratio features improved every model",
    "Linear RMSE fell 9.59 → 6.56 MPa",
    "Top predictors: log-age, water/binder, water/cement",
], size=17)

# 5 Results: leakage
s = prs.slides.add_slide(BLANK)
header(s, "Results", "Random splitting understates error by about a quarter")
s.shapes.add_picture("figures/fig2_split_leakage.png", Inches(0.6), Inches(1.8), height=Inches(5.2))
bullets(s, 7.9, 2.0, 5.0, [
    "Random split: 105 of 155 test mixtures also seen in training",
    "Mean RMSE 3.89 MPa (random) vs 5.12 MPa (grouped)",
    "≈5 MPa is the realistic error for a genuinely new mix",
], size=18)

# 6 Takeaway
s = prs.slides.add_slide(BLANK)
header(s, "Takeaway", "Feature design and evaluation design both matter")
bullets(s, 0.6, 1.9, 12, [
    "Tree ensembles with physically grounded ratio features predict new mixtures to about 4–5 MPa",
    "Learned importances agree with Abrams' law and hydration kinetics",
    "Accuracy claims from random splits on this dataset are optimistic; group by mixture",
    "Limits: laboratory data only, missing variables (cement type, curing conditions), point estimates",
    "Next: prediction intervals, monotonic boosting, validation on job-site data (Young et al., 2019)",
], size=21)

prs.save("slides.pptx")
print("slides.pptx written")
